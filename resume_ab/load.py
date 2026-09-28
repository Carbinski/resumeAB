"""Load resumes and job descriptions as plain text from .txt, .pdf, .docx, or .tex files."""

from __future__ import annotations

import hashlib
import re
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile

import docx
from docx.opc.exceptions import OpcError
from docx.table import Table
from pylatexenc.latex2text import LatexNodes2Text, MacroTextSpec
from pylatexenc.latex2text import get_default_latex_context_db as get_default_text_context_db
from pylatexenc.latexwalker import LatexWalkerError, get_default_latex_context_db
from pylatexenc.macrospec import MacroSpec
from pypdf import PdfReader
from pypdf.errors import PyPdfError


@dataclass(frozen=True)
class Document:
    text: str
    source_type: str


class LoadError(Exception):
    pass


class UnsupportedInputType(LoadError):
    pass


class UnreadableDocument(LoadError):
    pass


class EmptyDocument(LoadError):
    pass


def _normalize(text: str) -> str:
    lines = [re.sub(r"[^\S\n]+", " ", line).strip() for line in text.splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _read_plain(path: Path) -> str:
    return path.read_text()


def _read_pdf(path: Path) -> str:
    try:
        pages = [page.extract_text() for page in PdfReader(path).pages]
    except PyPdfError as error:
        raise UnreadableDocument(f"Could not read PDF {path}: {error}") from error
    return _normalize("\n".join(pages))


def _docx_table_rows(table: Table) -> Iterator[str]:
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            text = cell.text.strip()
            # python-docx repeats a merged cell once per grid column it spans.
            if text and (not cells or cells[-1] != text):
                cells.append(text)
        if cells:
            yield " | ".join(cells)


def _docx_blocks(container) -> Iterator[str]:
    for block in container.iter_inner_content():
        if isinstance(block, Table):
            yield from _docx_table_rows(block)
        else:
            yield block.text


def _read_docx(path: Path) -> str:
    try:
        document = docx.Document(str(path))
    except (OpcError, BadZipFile) as error:
        raise UnreadableDocument(f"Could not read Word document {path}: {error}") from error
    # Word resume templates often put the name and contact line in the page header.
    headers = [
        text
        for section in document.sections
        if not section.header.is_linked_to_previous
        for text in _docx_blocks(section.header)
    ]
    return _normalize("\n".join([*headers, *_docx_blocks(document)]))


def _href_to_text(node, l2tobj: LatexNodes2Text) -> str:
    # pylatexenc passes the converter by keyword, so the parameter must be named l2tobj.
    url = l2tobj.node_arg_to_text(node, 0).strip().removeprefix("mailto:")
    label = l2tobj.node_arg_to_text(node, 1).strip()
    return label if label == url else f"{label} <{url}>"


# pylatexenc's default parser gives \href no arguments, which crashes its own text renderer.
_LATEX_PARSE_CONTEXT = get_default_latex_context_db()
_LATEX_PARSE_CONTEXT.add_context_category(
    "resume_ab", prepend=True, macros=[MacroSpec("href", "{{")]
)
_LATEX_TEXT_CONTEXT = get_default_text_context_db()
_LATEX_TEXT_CONTEXT.add_context_category(
    "resume_ab", prepend=True, macros=[MacroTextSpec("href", _href_to_text)]
)


class _SpacedLatexToText(LatexNodes2Text):
    """Keeps arguments of unknown template macros like ``\\resumeSubheading{A}{B}`` apart."""

    def group_node_to_text(self, node) -> str:
        return f" {super().group_node_to_text(node)} "


def _read_latex(path: Path) -> str:
    source = path.read_text()
    begin, end = r"\begin{document}", r"\end{document}"
    if begin in source:
        source = source.split(begin, 1)[1].split(end, 1)[0]
    try:
        text = _SpacedLatexToText(latex_context=_LATEX_TEXT_CONTEXT).latex_to_text(
            source, latex_context=_LATEX_PARSE_CONTEXT
        )
    except LatexWalkerError as error:
        raise UnreadableDocument(f"Could not parse LaTeX {path}: {error}") from error
    return _normalize(text)


_FORMATS: dict[str, tuple[str, Callable[[Path], str]]] = {
    ".txt": ("plain", _read_plain),
    ".pdf": ("pdf", _read_pdf),
    ".docx": ("docx", _read_docx),
    ".tex": ("latex", _read_latex),
}


def _load_document(path: Path, kind: str) -> Document:
    suffix = path.suffix.lower()
    if suffix == ".doc":
        raise UnsupportedInputType(f"Legacy .doc {kind} files are not supported; save {path} as .docx")
    if suffix not in _FORMATS:
        raise UnsupportedInputType(
            f"Unsupported {kind} type: {path.suffix or '(none)'} (expected one of {', '.join(_FORMATS)})"
        )
    if not path.is_file():
        raise FileNotFoundError(f"{kind.capitalize()} file not found: {path}")

    source_type, extract = _FORMATS[suffix]
    text = extract(path)
    if not text.strip():
        hint = " Image-only (scanned) PDFs need OCR first." if source_type == "pdf" else ""
        raise EmptyDocument(f"No text found in {kind} {path}.{hint}")
    return Document(text, source_type)


def load_resumes(folder: Path) -> list[tuple[Path, Document]]:
    """Every supported resume in ``folder``, sorted by filename. Other files are skipped."""
    if not folder.is_dir():
        raise FileNotFoundError(f"Resume folder not found: {folder}")
    paths = sorted(
        path
        for path in folder.iterdir()
        if path.is_file() and path.suffix.lower() in _FORMATS and not path.name.startswith("~$")
    )
    return [(path, _load_document(path, "resume")) for path in paths]


def load_resume(path: Path) -> Document:
    return _load_document(path, "resume")


def load_job_description(source: Path | str) -> Document:
    if isinstance(source, str):
        return Document(source, "plain")
    return _load_document(source, "job description")


def content_sha256(text: str) -> str:
    """Identifies a document by content, so logs never need the text itself."""
    return hashlib.sha256(text.encode("utf-8")).hexdigest()
